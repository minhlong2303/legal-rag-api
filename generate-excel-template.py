#!/usr/bin/env python3
"""
Generate Excel template for Consultant Import
Creates consultant-import-template.xlsx with Vietnamese headers and sample data
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_consultant_template():
    """Create Excel template with Vietnamese headers"""

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Consultants"

    # Define headers in Vietnamese
    headers = [
        "Mã Tư Vấn",           # A - Consultant Code
        "Tên Tư Vấn",          # B - Consultant Name
        "Email",               # C - Email
        "Điện Thoại",          # D - Phone
        "Chuyên Môn",          # E - Specializations
        "Danh Mục",            # F - Categories
        "Năm Kinh Nghiệm",     # G - Experience Years
        "Đánh Giá TB (0-5)",   # H - Average Rating
        "Số Tư Vấn Hoàn Thành",# I - Completed Consultations
        "Thời Gian Phản Hồi (Phút)", # J - Response Time
        "Số Chat Tối Đa",      # K - Max Concurrent Chats
        "Website",             # L - Website
        "Địa Chỉ",             # M - Address
        "Thành Phố",           # N - City
        "Mô Tả",               # O - Description
        "Kích Hoạt"            # P - Is Active
    ]

    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header

        # Style header
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add border
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        cell.border = thin_border

    # Sample data (5 examples)
    sample_data = [
        {
            'code': 'CONS-001',
            'name': 'Nguyễn Văn Tư',
            'email': 'nguyen@consultant.com',
            'phone': '0123456789',
            'specializations': 'nhập khẩu hóa chất, xuất khẩu',
            'categories': 'import, chemical, permit',
            'experience': 15,
            'rating': 4.8,
            'consultations': 150,
            'response_time': 2,
            'max_chats': 5,
            'website': 'https://example.com',
            'address': '123 Đường Lê Lợi, Quận 1',
            'city': 'TP.HCM',
            'description': '15+ năm kinh nghiệm trong nhập khẩu hóa chất và xuất khẩu',
            'is_active': 'YES'
        },
        {
            'code': 'CONS-002',
            'name': 'Trần Thị Tư Vấn',
            'email': 'tran@consultant.com',
            'phone': '0987654321',
            'specializations': 'hóa chất, giấy phép môi trường',
            'categories': 'chemical, permit, environment',
            'experience': 10,
            'rating': 4.5,
            'consultations': 120,
            'response_time': 3,
            'max_chats': 4,
            'website': 'https://tran-consulting.vn',
            'address': '456 Đường Nguyễn Huệ, Quận 1',
            'city': 'TP.HCM',
            'description': '10+ năm tư vấn về giấy phép và quy định môi trường',
            'is_active': 'YES'
        },
        {
            'code': 'CONS-003',
            'name': 'Lê Văn Bình',
            'email': 'le.van.binh@consultant.vn',
            'phone': '0918765432',
            'specializations': 'giấy phép, pháp luật môi trường',
            'categories': 'permit, environment, legal',
            'experience': 12,
            'rating': 4.7,
            'consultations': 98,
            'response_time': 5,
            'max_chats': 3,
            'website': '',
            'address': '789 Đường D, Quận 2',
            'city': 'TP.HCM',
            'description': 'Chuyên gia pháp luật môi trường hàng đầu',
            'is_active': 'YES'
        },
        {
            'code': 'CONS-004',
            'name': 'Phạm Thị Dung',
            'email': 'pham.dung@consultant.com',
            'phone': '0945123456',
            'specializations': 'nhập khẩu, tuân thủ quy định',
            'categories': 'import, compliance, chemical',
            'experience': 8,
            'rating': 4.3,
            'consultations': 67,
            'response_time': 4,
            'max_chats': 3,
            'website': 'https://phampham.com',
            'address': '101 Đường Pasteur, Quận 3',
            'city': 'TP.HCM',
            'description': '8 năm kinh nghiệm tư vấn nhập khẩu',
            'is_active': 'YES'
        },
        {
            'code': 'CONS-005',
            'name': 'Đỗ Văn Minh',
            'email': 'do.minh@expert.vn',
            'phone': '0912345678',
            'specializations': 'hóa chất nguy hại, xử lý chất thải',
            'categories': 'chemical, waste, hazmat',
            'experience': 18,
            'rating': 4.9,
            'consultations': 200,
            'response_time': 1,
            'max_chats': 6,
            'website': 'https://dominh-experts.com',
            'address': '999 Đường Cách Mạng Tháng Tám, Quận 3',
            'city': 'TP.HCM',
            'description': '18 năm chuyên gia hóa chất nguy hại',
            'is_active': 'YES'
        }
    ]

    # Write sample data
    for row_num, data in enumerate(sample_data, 2):
        data_row = [
            data['code'],
            data['name'],
            data['email'],
            data['phone'],
            data['specializations'],
            data['categories'],
            data['experience'],
            data['rating'],
            data['consultations'],
            data['response_time'],
            data['max_chats'],
            data['website'],
            data['address'],
            data['city'],
            data['description'],
            data['is_active']
        ]

        for col_num, value in enumerate(data_row, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value

            # Style data cells
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            cell.border = thin_border

            # Align
            if col_num in [7, 8, 9, 10, 11]:  # Number columns
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            # Format specific columns
            if col_num == 8:  # Rating - number format
                cell.number_format = '0.0'

    # Set column widths
    column_widths = [15, 18, 22, 15, 30, 25, 15, 12, 18, 15, 12, 20, 25, 12, 40, 12]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    # Set row height for header
    ws.row_dimensions[1].height = 30

    # Freeze header row
    ws.freeze_panes = "A2"

    # Save file
    file_path = "consultant-import-template.xlsx"
    wb.save(file_path)
    print(f"✅ Template created successfully: {file_path}")
    print(f"📊 Columns: {len(headers)}")
    print(f"📝 Sample rows: {len(sample_data)}")
    print("\n📋 Column Names (Tiếng Việt):")
    for col_num, header in enumerate(headers, 1):
        print(f"   {get_column_letter(col_num)}: {header}")

if __name__ == "__main__":
    create_consultant_template()

